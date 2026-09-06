"""Извлечение текста из резюме и его структурирование."""

from __future__ import annotations

import asyncio
import io
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, List, Optional, Tuple

from src.models.models import ResumeProfile
from src.utils.llm import OpenRouterAuthError, create_chat_model
from src.utils.logger import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".txt", ".md", ".rtf"})

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_PDF_PAGES = 30
MAX_LLM_CHARS = 16000

# Ниже этого объёма текста PDF считается сканом (изображением)
MIN_TEXT_PER_PAGE = 80

_LIGATURES = {
    "ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl",
    "ﬅ": "st", "ﬆ": "st",
    "•": "· ", "\uf0b7": "· ", "\uf0a7": "· ", "\u2022": "· ",
    "\u00a0": " ", "\u200b": "", "\ufeff": "", "\u2013": "-", "\u2014": "—",
}

RESUME_ANALYSIS_PROMPT = """
Ты — HR-аналитик. Проанализируй резюме кандидата и извлеки структурированный профиль для поиска вакансий.

Текст резюме:
{text}

Инструкции:
1. desired_role — целевая должность (на русском или английском, как в резюме).
2. skills — до 20 ключевых навыков (технологии, инструменты, soft skills).
3. experience_years — общий стаж ТОЛЬКО числом (например 6 или 6.5), без «+» и текста.
4. experience_summary — 2–4 предложения о карьере и ключевых достижениях.
5. seniority_level — junior / mid / senior / lead / principal (по стажу и ролям).
6. work_history — до 5 записей «Компания — должность (годы)».
7. languages — языки с уровнем, например «English B2», «Polski A2».
8. industries — отрасли: fintech, e-commerce, healthcare и т.д.
9. location — город и страна, где кандидат ищет работу. Бери город проживания
   или явно указанное предпочтение. Формат: «Warszawa, Polska». Если города
   в резюме нет — оставь null, не угадывай по языку резюме.
10. salary_expectation — ожидания по зарплате, если есть.
11. summary — краткий профиль кандидата в 1–2 предложения для поискового запроса.

Если данных нет — оставь поле пустым или null, не выдумывай.
"""


class ResumeExtractionError(Exception):
    """Текст из файла резюме получить не удалось."""


@dataclass
class ExtractedDocument:
    text: str
    engine: str
    pages: int = 0
    is_scanned: bool = False


# ─────────────────────────── очистка текста ───────────────────────────

def clean_resume_text(raw: str) -> str:
    """Приводит выдачу PDF-парсера к читаемому виду."""
    if not raw:
        return ""

    text = unicodedata.normalize("NFKC", raw)
    for source, target in _LIGATURES.items():
        text = text.replace(source, target)

    # Управляющие символы, кроме переводов строк и табуляции
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text)

    # Переносы слов на границе строк: «разработ-\nчик» → «разработчик»
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)

    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    lines = _drop_repeated_lines(lines)

    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


_FOOTER_PATTERN = re.compile(
    r"^(?:strona|page|стр\.?|страница)?\s*\d+\s*(?:/|z|of|из)?\s*\d*$"
    r"|^https?://\S+$"
    r"|(?:^|\s)(?:curriculum vitae|cv\b)",
    re.I,
)


def _drop_repeated_lines(lines: List[str]) -> List[str]:
    """
    Убирает колонтитулы и номера страниц.

    Осторожно: повторяющиеся строки резюме («Обязанности:») удалять нельзя,
    поэтому отбрасываются только строки, похожие на служебные.
    """
    counts = Counter(line for line in lines if line)
    noisy = {
        line
        for line, count in counts.items()
        if count >= 3 and len(line) <= 60 and _FOOTER_PATTERN.search(line)
    }
    noisy.update(line for line in counts if line and _FOOTER_PATTERN.fullmatch(line))
    if not noisy:
        return lines
    return [line for line in lines if line not in noisy]


def _text_density(text: str, pages: int) -> float:
    return len(text.strip()) / max(pages, 1)


# ─────────────────────────── PDF ───────────────────────────

def _extract_pdf_pypdf(data: bytes) -> ExtractedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))

    if reader.is_encrypted:
        # Большинство «защищённых» резюме зашифрованы пустым паролем
        for password in ("", " "):
            try:
                if reader.decrypt(password):
                    break
            except Exception:
                continue
        else:
            raise ResumeExtractionError(
                "PDF защищён паролем. Снимите пароль или отправьте DOCX/TXT."
            )

    pages = reader.pages[:MAX_PDF_PAGES]
    parts = []
    for page in pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception as error:
            logger.debug("pypdf не смог прочитать страницу: %s", error)

    text = "\n".join(parts)
    return ExtractedDocument(text=text, engine="pypdf", pages=len(pages))


def _extract_pdf_pdfplumber(data: bytes) -> ExtractedDocument:
    import pdfplumber

    parts: List[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = pdf.pages[:MAX_PDF_PAGES]
        for page in pages:
            parts.append(page.extract_text(layout=False) or "")
            for table in page.extract_tables() or []:
                for row in table:
                    cells = [str(c).strip() for c in row if c and str(c).strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        page_count = len(pages)

    return ExtractedDocument(text="\n".join(parts), engine="pdfplumber", pages=page_count)


def _extract_pdf_pymupdf(data: bytes) -> ExtractedDocument:
    import fitz  # PyMuPDF

    parts: List[str] = []
    with fitz.open(stream=data, filetype="pdf") as document:
        page_count = min(document.page_count, MAX_PDF_PAGES)
        for index in range(page_count):
            parts.append(document.load_page(index).get_text("text"))

    return ExtractedDocument(text="\n".join(parts), engine="pymupdf", pages=page_count)


def _ocr_pdf(data: bytes) -> ExtractedDocument:
    """OCR для сканов. Работает только если установлены pytesseract и Tesseract."""
    import pytesseract
    from pdf2image import convert_from_bytes

    images = convert_from_bytes(data, dpi=200, fmt="png")[:10]
    parts = [
        pytesseract.image_to_string(image, lang="pol+eng+rus")
        for image in images
    ]
    return ExtractedDocument(
        text="\n".join(parts), engine="ocr", pages=len(images), is_scanned=True,
    )


def extract_text_from_pdf_sync(data: bytes) -> ExtractedDocument:
    """
    Пробует несколько движков и берёт лучший результат.

    Разные PDF (текстовый слой, таблицы, вёрстка в колонки) читаются
    разными библиотеками по-разному, поэтому выбор делается по объёму текста.
    """
    engines: Tuple[Tuple[str, Callable[[bytes], ExtractedDocument]], ...] = (
        ("pypdf", _extract_pdf_pypdf),
        ("pdfplumber", _extract_pdf_pdfplumber),
        ("pymupdf", _extract_pdf_pymupdf),
    )

    best: Optional[ExtractedDocument] = None
    errors: List[str] = []

    for name, extractor in engines:
        try:
            document = extractor(data)
        except ResumeExtractionError:
            raise
        except ImportError:
            logger.debug("Движок %s не установлен", name)
            continue
        except Exception as error:
            errors.append(f"{name}: {error}")
            logger.debug("Движок %s не справился: %s", name, error)
            continue

        document = ExtractedDocument(
            text=clean_resume_text(document.text),
            engine=document.engine,
            pages=document.pages,
        )
        if best is None or len(document.text) > len(best.text):
            best = document

        # Достаточно плотного текста — дальше не перебираем
        if _text_density(document.text, document.pages) >= 400:
            break

    if best is None:
        detail = "; ".join(errors) or "неизвестная ошибка"
        raise ResumeExtractionError(f"Не удалось прочитать PDF ({detail}).")

    if _text_density(best.text, best.pages) < MIN_TEXT_PER_PAGE:
        logger.info("PDF похож на скан (%d симв. на %d стр.) — пробуем OCR",
                    len(best.text), best.pages)
        try:
            ocr = _ocr_pdf(data)
            ocr.text = clean_resume_text(ocr.text)
            if len(ocr.text) > len(best.text):
                return ocr
        except ImportError:
            best.is_scanned = True
        except Exception as error:
            logger.warning("OCR не сработал: %s", error)
            best.is_scanned = True

    return best


async def extract_text_from_pdf(data: bytes) -> str:
    document = await asyncio.to_thread(extract_text_from_pdf_sync, data)
    if document.is_scanned and len(document.text) < 200:
        raise ResumeExtractionError(
            "Похоже, это скан или фото без текстового слоя.\n"
            "Отправьте PDF с выделяемым текстом, DOCX или пришлите резюме текстом."
        )
    logger.info("PDF прочитан движком %s: %d симв., %d стр.",
                document.engine, len(document.text), document.pages)
    return document.text


# ─────────────────────────── DOCX и текст ───────────────────────────

def extract_text_from_docx_sync(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts: List[str] = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))

    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                if paragraph.text.strip():
                    parts.append(paragraph.text.strip())

    return clean_resume_text("\n".join(parts))


async def extract_text_from_docx(data: bytes) -> str:
    return await asyncio.to_thread(extract_text_from_docx_sync, data)


def _strip_rtf(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    text = re.sub(r"\\'([0-9a-fA-F]{2})", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return clean_resume_text(text)


async def extract_text_from_file(filename: str, data: bytes) -> str:
    """Извлекает текст из PDF, DOCX, RTF или plain text."""
    if not data:
        raise ResumeExtractionError("Файл пустой.")
    if len(data) > MAX_FILE_BYTES:
        raise ResumeExtractionError(
            f"Файл больше {MAX_FILE_BYTES // (1024 * 1024)} МБ. Пришлите версию поменьше."
        )

    extension = Path(filename or "").suffix.lower()

    if extension == ".pdf" or data[:5] == b"%PDF-":
        return await extract_text_from_pdf(data)
    if extension == ".docx" or data[:2] == b"PK":
        return await extract_text_from_docx(data)
    if extension == ".doc":
        raise ResumeExtractionError(
            "Старый формат .doc не поддерживается.\n"
            "Сохраните резюме как PDF или DOCX и отправьте снова."
        )
    if extension == ".rtf":
        return _strip_rtf(data)
    if extension in {".txt", ".md"}:
        return clean_resume_text(data.decode("utf-8", errors="replace"))

    raise ResumeExtractionError(
        f"Формат {extension or 'файла'} не поддерживается. "
        f"Используйте: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
    )


# ─────────────────────────── анализ через LLM ───────────────────────────

async def analyze_resume_text(text: str) -> ResumeProfile:
    """Структурирует текст резюме через LLM."""
    llm = create_chat_model()
    structured_llm = llm.with_structured_output(ResumeProfile)

    prompt = RESUME_ANALYSIS_PROMPT.format(text=text[:MAX_LLM_CHARS])
    result: ResumeProfile = await asyncio.to_thread(structured_llm.invoke, prompt)
    result = result.model_copy(update={"raw_text": text[:MAX_LLM_CHARS]})
    logger.info(
        "Резюме проанализировано: %s, опыт %s лет, %d навыков, локация %s",
        result.desired_role, result.experience_years,
        len(result.skills), result.location or "—",
    )
    return result


async def analyze_resume_file(filename: str, data: bytes) -> ResumeProfile:
    """Извлекает текст из файла и анализирует резюме."""
    text = await extract_text_from_file(filename, data)
    if len(text.strip()) < 80:
        raise ResumeExtractionError(
            "В файле почти нет текста. Проверьте, что резюме не является "
            "картинкой, и попробуйте PDF с выделяемым текстом или DOCX."
        )
    return await analyze_resume_text(text)


def profile_to_search_query(profile: ResumeProfile) -> str:
    """Формирует поисковый запрос из профиля резюме."""
    parts = [profile.desired_role]
    if profile.seniority_level:
        parts.append(profile.seniority_level)
    if profile.skills:
        parts.append(", ".join(profile.skills[:8]))
    if profile.industries:
        parts.append(", ".join(profile.industries[:3]))
    if profile.summary:
        parts.append(profile.summary[:180])
    return " ".join(p for p in parts if p)


def format_llm_error(error: Exception) -> str:
    """Понятное сообщение об ошибке для пользователя."""
    if isinstance(error, (ResumeExtractionError, ValueError)):
        return f"❌ {error}"
    if isinstance(error, OpenRouterAuthError):
        return str(error)

    text = str(error).lower()
    if "401" in text or "user not found" in text:
        return (
            "❌ Ключ OpenRouter недействителен (401).\n\n"
            "1. Откройте https://openrouter.ai/keys\n"
            "2. Создайте новый API key\n"
            "3. Вставьте в .env: OPENROUTER_API_KEY=sk-or-v1-...\n"
            "4. Перезапустите бота"
        )
    if "402" in text or "insufficient" in text or "credit" in text:
        return (
            "❌ Недостаточно кредитов OpenRouter.\n"
            "Пополните баланс: https://openrouter.ai/credits "
            "или задайте бесплатную модель: OPENROUTER_MODEL=google/gemini-2.0-flash-exp:free"
        )
    return f"❌ Ошибка AI: {error}"
